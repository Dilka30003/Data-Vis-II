"""Reproducible extraction. Run with Python 3 + openpyxl + shapely.

python scripts/extract_data.py --raw ../source-data --out data
No API keys, imputation or simulated observations.
"""
import argparse
import csv
import hashlib
import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
args = argparse.ArgumentParser()
args.add_argument("--raw", type=Path, default=ROOT.parent / "source-data")
args.add_argument("--out", type=Path, default=ROOT / "data")
opts = args.parse_args()
OUT = opts.out
OUT.mkdir(parents=True, exist_ok=True)
SOURCES = [
    dict(source_id="acma_internet", provider="ACMA", filename="How we use the internet data tables.xlsx", url="https://www.acma.gov.au/sites/default/files/2026-02/How%20we%20use%20the%20internet%20data%20tables.xlsx", report_url="https://www.acma.gov.au/publications/2026-02/report/communications-and-media-australia-how-we-use-internet"),
    dict(source_id="acma_communication", provider="ACMA", filename="How we communicate data tables.xlsx", url="https://www.acma.gov.au/sites/default/files/2026-02/How%20we%20communicate%20data%20tables.xlsx", report_url="https://www.acma.gov.au/publications/2026-02/report/communications-and-media-australia-how-we-communicate"),
    dict(source_id="accc", provider="ACCC", filename="mobile-infrastructure-report-2026-output-tables.xlsx", url="https://www.accc.gov.au/system/files/mobile-infrastructure-report-2026-output-tables.xlsx", report_url="https://www.accc.gov.au/by-industry/telecommunications-and-internet/mobile-services-regulation/mobile-infrastructure-report/mobile-infrastructure-report-2026"),
    dict(source_id="accc_sites", provider="ACCC", filename="mobile-sites-telstra-2026.csv", url="https://data.gov.au/data/dataset/4b472a18-d0fa-409c-994a-ab17162bcb90/resource/b508e517-a2d1-41aa-a8f4-3ec7e57e0cb3/download/mobile-sites-telstra-2026.csv", report_url="https://data.gov.au/data/dataset/accc-mobile-infrastructure-report-data-release/resource/b508e517-a2d1-41aa-a8f4-3ec7e57e0cb3"),
    dict(source_id="adii", provider="ADII", filename="adii_report.pdf", url="https://digitalinclusionindex.org.au/wp-content/uploads/2025/10/ADII-Report-2025_V6-Remediated.pdf", report_url="https://digitalinclusionindex.org.au/download-reports/"),
    dict(source_id="natural_earth", provider="Natural Earth", filename="ne_states.geojson", url="https://raw.githubusercontent.com/nvkelso/natural-earth-vector/master/geojson/ne_50m_admin_1_states_provinces.geojson", report_url="https://www.naturalearthdata.com/about/terms-of-use/")
]
MANIFEST = []


def write_csv(name, records, description):
    assert records, name
    keys = list(records[0])
    with (OUT / name).open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=keys)
        w.writeheader()
        w.writerows(records)
    MANIFEST.append(dict(file=name, rows=len(records), columns=keys, description=description))


for source in SOURCES:
    path = opts.raw / source["filename"]
    source["sha256"] = hashlib.sha256(path.read_bytes()).hexdigest() if path.exists() else ""
    source["retrieved"] = "2026-10-17"
write_csv("sources.csv", SOURCES, "Original sources; join other CSVs on source_id for the full source URL.")

iwb = openpyxl.load_workbook(opts.raw / SOURCES[0]["filename"], data_only=True)
cwb = openpyxl.load_workbook(opts.raw / SOURCES[1]["filename"], data_only=True)
awb = openpyxl.load_workbook(opts.raw / SOURCES[2]["filename"], data_only=True)


def year_columns(ws):
    for row in ws.iter_rows():
        if sum(str(c.value).isdigit() and 2017 <= int(str(c.value)) <= 2026 for c in row) >= 4:
            return row[0].row, [c.column for c in row if str(c.value).isdigit() and 2017 <= int(str(c.value)) <= 2026]
    raise ValueError(ws.title)


def labels(ws, header):
    # Six-row statistic blocks: %, response n, base n, weighted n,
    # weighted base, comparison symbols. Do not treat count rows as percents.
    found = []
    for r in range(header + 1, ws.max_row - 4):
        label = ws.cell(r, 1).value
        if label and isinstance(label, str) and not label.startswith(("Column", "Filter:", "Multiple comparison", "Back to", "Weight:")):
            if isinstance(ws.cell(r, 2).value, (int, float)):
                found.append((r, label))
    return found


def live_label(label):
    return not any(s in label for s in ["UP TO ", " ONLY)", "Invalid", "BACKCODED"])


def acma(wb, sheet, sid, field, demographic=False, current_only=False):
    ws = wb[sheet]
    hr, cols = year_columns(ws)
    recs = []
    groups = {}
    if demographic:
        group, dim = "", ""
        for c in range(2, ws.max_column + 1):
            heading = ws.cell(hr - 3, c).value
            if heading:
                dim = "age" if str(heading).startswith("Age") else "gender" if str(heading).startswith("Gender") else "location"
            if ws.cell(hr - 2, c).value:
                group = ws.cell(hr - 2, c).value
            groups[c] = (dim, group)
    for r, label in labels(ws, hr):
        if not live_label(label):
            continue
        for c in cols:
            year = int(ws.cell(hr, c).value)
            if current_only and year != 2025:
                continue
            value = ws.cell(r, c).value
            if not isinstance(value, (float, int)):
                continue
            status = "published"
            # Historic zeros are ambiguous in these exports; do not turn
            # a not-yet-asked option into a real zero observation.
            if year < 2025 and value == 0:
                status = "zero_or_not_asked"
            base = ws.cell(r + 2, c).value
            rec = dict(year=year, **{field: label}, percent=round(value * 100, 6) if status == "published" else "", base_n=base, response_n=ws.cell(r+1, c).value, status=status, small_base=1 if isinstance(base, (int,float)) and base < 100 else 0)
            if demographic:
                rec.update(dimension=groups[c][0], group=groups[c][1])
            rec.update(source_id=sid, source_sheet=sheet, source_cell=ws.cell(r,c).coordinate)
            recs.append(rec)
    return recs


write_csv("device_trends.csv", acma(iwb, "QD5. Devices  by Study year", "acma_internet", "device"), "Weighted % of adults using each device for internet access in past six months. 2017–2025. Multiple response.")
write_csv("device_demographics_2025.csv", acma(iwb, "QD5. Devices  by PBI BANNER", "acma_internet", "device", True, True), "Device use by age, gender and metro/regional location, 2025; all-adult bases per group.")
write_csv("phone_age_trends.csv", [r for r in acma(iwb, "QD5. Devices  by PBI BANNER", "acma_internet", "device", True) if r["device"]=="Mobile phone" and r["dimension"]=="age"], "Mobile-phone internet use by age, 2017–2025. Weighted adult percentages; small-base flags retained.")
write_csv("activity_2025.csv", acma(iwb, "QD8. Internet by Study year", "acma_internet", "activity", current_only=True), "All-device online activities, six-month recall, 2025. Not phone-only. Multiple response.")
write_csv("communication_trends.csv", acma(cwb, "QA1. In the p by Study year", "acma_communication", "service"), "Six-month personal communication-service use, 2017–2025. NET categories overlap their components.")
write_csv("communication_demographics_2025.csv", acma(cwb, "QA1. In the p by PBI BANNER", "acma_communication", "service", True, True), "2025 communication-service use by demographic group.")
write_csv("platform_trends.csv", acma(cwb, "QE2. Which we by Study year", "acma_communication", "platform"), "All-device six-month website/app use, 2021–2025. Not smartphone-only.")
write_csv("platform_demographics_2025.csv", acma(cwb, "QE2. Which we by PBI BANNER", "acma_communication", "platform", True, True), "2025 six-month website/app use by age/gender/location. Preserve this survey's base, not the device survey's.")
write_csv("platform_weekly_2025.csv", acma(cwb, "QE12. Which o by Study year", "acma_communication", "platform", current_only=True), "Past-seven-day platform use, 2025. Do not conflate with six-month prevalence.")

# QD6: device-user denominator, not all adults.
ws = iwb["QD6. How ofte by Study year"]
freqs = []
frequency = None
for c in range(2, ws.max_column+1):
    if ws.cell(3,c).value:
        frequency = ws.cell(3,c).value
    if not live_label(str(frequency)) or frequency in ["NET", "Multiple times a day + Once a day", "Skipped due to backcoding"]:
        continue
    year = ws.cell(4,c).value
    if not str(year).isdigit():
        continue
    for r in range(5, 104, 7):
        device, val = ws.cell(r,1).value, ws.cell(r,c).value
        if not device or not live_label(device) or not isinstance(val, (int,float)):
            continue
        freqs.append(dict(year=int(year), device=device, frequency=frequency, percent=round(val*100,6), base_n=ws.cell(r+2,c).value, response_n=ws.cell(r+1,c).value, source_id="acma_internet", source_sheet=ws.title, source_cell=ws.cell(r,c).coordinate))
write_csv("device_frequency.csv", freqs, "Frequency within users of each device. Exclusive frequency bins; summed daily NET omitted.")

# Latest QD6 age export is ROW %, not P(frequency | age).
# Derive P(frequency | age) from weighted count / weighted NET age count.
ws = iwb["QD6. How ofte by Age (consolida"]
agefreq = []
frequency = None
for c in range(2, 90):
    if ws.cell(3,c).value:
        frequency = ws.cell(3,c).value
    age = ws.cell(4,c).value
    if age == "NET" or not live_label(str(frequency)) or frequency == "Skipped due to backcoding":
        continue
    age_index = (c-2) % 8
    den_col = 98 + age_index  # CT: weighted NET for each age, same Mobile phone row
    den = ws.cell(7,den_col).value
    num = ws.cell(7,c).value
    if isinstance(num,(float,int)) and isinstance(den,(float,int)) and den:
        agefreq.append(dict(year=2025, device="Mobile phone", age_group=age, frequency=frequency, percent=round(100*num/den,6), base_n=ws.cell(6,den_col).value, response_n=ws.cell(6,c).value, weighted_count=num, weighted_base=den, source_id="acma_internet", source_sheet=ws.title, source_cell=f"{ws.cell(7,c).coordinate}/{ws.cell(7,den_col).coordinate}", method="weighted_count / weighted_age_NET * 100"))
write_csv("phone_frequency_by_age_2025.csv", agefreq, "Derived within-age frequency percentages from weighted QD6 counts; NOT the source's row percentages.")


def accc_table(number):
    for sheet in awb:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value,str) and re.match(rf"Table {number}:",cell.value):
                    r,c = cell.row,cell.column
                    ycols = []
                    for cc in range(c+1,sheet.max_column+1):
                        value = str(sheet.cell(r+1,cc).value)
                        if value.isdigit() and 2018 <= int(value) <= 2026:
                            ycols.append((cc,int(value)))
                        else:
                            break
                    records,group = [],"Australia"
                    for rr in range(r+2,sheet.max_row+1):
                        name=sheet.cell(rr,c).value
                        if isinstance(name,str) and name.startswith("Table "):
                            break
                        if not name:
                            continue
                        if name not in ["Optus","Telstra","TPG","Optus-TPG MOCN"]:
                            group=str(name)
                            continue
                        for cc,year in ycols:
                            v=sheet.cell(rr,cc).value
                            if isinstance(v,(int,float)):
                                records.append(dict(year=year,operator=name,geography=group,count=int(v),table=number,source_sheet=sheet.title,source_cell=sheet.cell(rr,cc).coordinate))
                    return records
    raise ValueError(number)


states_names = dict(NSW="New South Wales", VIC="Victoria", QLD="Queensland", SA="South Australia", WA="Western Australia", TAS="Tasmania", NT="Northern Territory", ACT="Australian Capital Territory")
# Cartographic label positions, not source observations.
anchors = dict(NSW=(147.0,-32.3), VIC=(144.7,-37.2), QLD=(144.6,-23.0), SA=(135.0,-30.2), WA=(122.0,-25.0), TAS=(146.6,-42.0), NT=(133.4,-19.1), ACT=(149.13,-35.35))


def joined(tables, geo_field):
    output={}
    for table,field in tables:
        for rec in accc_table(table):
            k=(rec["year"],rec["operator"],rec["geography"])
            if k not in output:
                output[k]=dict(year=k[0],reference_date=f"{k[0]}-01-31",operator=k[1],**{geo_field:k[2]},source_id="accc")
            output[k][field]=rec["count"]
            output[k][field+"_cell"]=f"{rec['source_sheet']}!{rec['source_cell']} (Table {table})"
    rows=list(output.values())
    allfields=sorted(set().union(*(x.keys() for x in rows)))
    result=[]
    for rec in rows:
        den=rec.get("total_sites")
        rec["five_g_site_share"]=round(100*rec["five_g_sites"]/den,6) if den and rec.get("five_g_sites") is not None else ""
        if geo_field == "state":
            rec.update(state_name=states_names.get(rec["state"],rec["state"]),longitude=anchors.get(rec["state"],("",""))[0],latitude=anchors.get(rec["state"],("",""))[1])
        result.append(rec)
    keys=list(dict.fromkeys(k for rec in result for k in rec))
    return [{k:rec.get(k,"") for k in keys} for rec in result]


write_csv("network_states.csv", joined([(3,"total_sites"),(7,"five_g_sites"),(35,"new_sites"),(42,"black_spot_sites")],"state"), "Operator/state counts and derived 5G-site shares; 31 January observations, 2018–2026. Not coverage.")
write_csv("network_remoteness.csv", joined([(2,"total_sites"),(6,"five_g_sites"),(34,"new_sites")],"remoteness"), "Operator/remoteness counts and derived 5G-site shares. ABS remoteness classification break in 2023.")
write_csv("network_national.csv", joined([(1,"total_sites"),(5,"five_g_sites"),(33,"new_sites"),(40,"black_spot_sites")],"geography"), "National operator totals. MOCN remains separate; do not sum to unique physical towers.")

# Published ADII figures transcribed from verified report text, never invented
# microdata. Data collected 2024; report edition 2025.
adii = []
for group,value in [("All Australians",9.7),("Outside capital cities",12.4),("People with disability",14.3),("Did not complete secondary school",17.0),("Public housing residents",18.8)]:
    adii.append(dict(data_year=2024,report_year=2025,measure="Mobile-only internet access",group=group,percent=value,source_id="adii",source_location="Access section: Mobile-only connections; published group percentages",method="Published-value transcription; groups overlap"))
adii.append(dict(data_year=2024,report_year=2025,measure="Smartphone ownership",group="All Australians",percent=93.3,source_id="adii",source_location="Access section: Not everyone has access to appropriate devices",method="Published-value transcription"))
write_csv("adii_mobile_access.csv",adii,"Published mobile-only access and smartphone ownership percentages. Groups overlap; no respondent microdata.")
adii_states=[]
for state,vals in dict(NSW=(74.0,77.0,71.1,73.8),VIC=(74.1,77.5,70.5,74.3),QLD=(72.9,76.3,69.3,73.0),SA=(71.3,75.7,66.9,71.3),WA=(73.9,77.1,70.8,73.7),TAS=(69.4,73.0,65.4,69.7),NT=(72.5,71.7,71.9,73.7),ACT=(79.5,80.9,78.0,79.6)).items():
    for measure,value in zip(["Index","Access","Affordability","Digital ability"],vals):
        adii_states.append(dict(data_year=2024,report_year=2025,state=state,state_name=states_names[state],measure=measure,score=value,source_id="adii",source_location="Figure 1: 2025 National ADII scores",method="Published-value transcription; small area estimates"))
write_csv("adii_state_scores.csv",adii_states,"ADII Figure 1 state/territory scores out of 100. NOT smartphone-use percentages.")

# Public-domain Natural Earth generalised states; used for illustrative mapping,
# not statistical geographic classification or official survey boundaries.
from shapely.geometry import shape, mapping, Point
geo=json.loads((opts.raw/"ne_states.geojson").read_text())
features=[]
for f in geo["features"]:
    p=f["properties"]
    if p.get("admin") != "Australia":
        continue
    name=p.get("name")
    state=next((s for s,n in states_names.items() if name==n), None)
    if state:
        geom=shape(f["geometry"]).simplify(.018,preserve_topology=True)
        features.append(dict(type="Feature",properties=dict(state=state,state_name=states_names[state]),geometry=mapping(geom)))
assert len(features)==8, [(f["properties"]) for f in features]
(OUT/"australian_states.geojson").write_text(json.dumps(dict(type="FeatureCollection",features=features),separators=(",",":")))

victoria=shape(next(f["geometry"] for f in features if f["properties"]["state"]=="VIC"))
sites=[]
with (opts.raw/"mobile-sites-telstra-2026.csv").open(newline="",encoding="utf-8-sig") as f:
    seen=set()
    for row_num,row in enumerate(csv.DictReader(f),2):
        lat,lon=float(row["Latitude"]),float(row["Longitude"])
        site_id=row["RFNSA ID"].strip()
        key=(site_id,lat,lon)
        seen.add(key)
        if not (-90 <= lat <=90 and -180<=lon<=180):
            raise ValueError(key)
        nr=any(str(v).strip().upper()=="Y" for k,v in row.items() if k.startswith("NR"))
        sites.append(dict(year=int(row["Year"]),operator=row["MNO"],site_id=site_id,latitude=lat,longitude=lon,five_g_present=1 if nr else 0,technology="5G present" if nr else "No 5G band recorded",co_funded=row["Co_funded"] or "Not recorded",co_funding_program=row["Co_contribution_program"] or "Not recorded",source_id="accc_sites",source_row=row_num))
from collections import Counter
location_counts=Counter((s["site_id"],s["latitude"],s["longitude"]) for s in sites)
for s in sites:
    s["records_at_same_id_coordinates"]=location_counts[(s["site_id"],s["latitude"],s["longitude"])]
write_csv("telstra_sites_2026.csv",sites,"Complete current Telstra site records, coordinate-checked; repeated IDs/coordinates intentionally retained (bands may differ). source_row is the unique row key. NR radio-band Y means 5G present. Missing co-funding is not imputed to No.")
vic_sites=[s for s in sites if victoria.covers(Point(s["longitude"],s["latitude"]))]
write_csv("telstra_sites_victoria_2026.csv",vic_sites,"Subset intersecting generalised Natural Earth Victoria polygon; illustrative, not an official state classification. Border/coastal sites can differ.")

# Data QA: provenance, bounds, current bases, and independently published totals.
for entry in MANIFEST:
    with (OUT/entry["file"]).open(newline="",encoding="utf-8") as f:
        for row in csv.DictReader(f):
            if row.get("percent") not in (None,""):
                assert 0<=float(row["percent"])<=100,(entry["file"],row)
assert len(sites)==12175, f"Telstra CSV count {len(sites)} differs from ACCC national Table 1: 12175"
state_rows=joined([(3,"total_sites"),(7,"five_g_sites")],"state")
for op in ["Optus","Telstra","TPG"]:
    latest=[r for r in state_rows if r["operator"]==op and r["year"]==2026]
    national=[r for r in joined([(1,"total_sites"),(5,"five_g_sites")],"geography") if r["operator"]==op and r["year"]==2026][0]
    assert sum(r["total_sites"] for r in latest)==national["total_sites"],op
    assert sum(r["five_g_sites"] for r in latest if isinstance(r["five_g_sites"],(int,float)))==national["five_g_sites"],op
for age in {r["age_group"] for r in agefreq}:
    total=sum(r["percent"] for r in agefreq if r["age_group"]==age)
    assert abs(total-100)<.01,(age,total)
for entry in MANIFEST:
    entry["bytes"]=(OUT/entry["file"]).stat().st_size
summary=dict(extracted="2026-10-17",percent_scale="0–100, not 0–1",missing_values="Empty numeric fields are missing, never zero-filled.",files=MANIFEST,validation=dict(telstra_site_record_count=len(sites),distinct_id_coordinate_locations=len(location_counts),victoria_subset=len(vic_sites),state_totals_match_national=True,frequency_age_sums_match_100=True,percentage_bounds_pass=True))
(OUT/"manifest.json").write_text(json.dumps(summary,indent=2))
print(json.dumps(summary,indent=2))
